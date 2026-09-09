"""Tests for document creation helpers."""

from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from PIL import Image

from pavlovia_sign_report.document_creator import (
    _decode_image,
    create_doc_with_table,
    create_single_doc,
    render_template,
    render_value,
    validate_image_value,
)
from pavlovia_sign_report.models import ColumnSettings, ColumnType, DocumentSettings


def test_decode_image_strips_prefix(png_data_uri: str):
    image = Image.open(_decode_image(png_data_uri))
    assert image.size == (10, 10)


def test_decode_image_rejects_invalid_base64():
    with pytest.raises(ValueError, match="valid base64"):
        _decode_image("data:image/png;base64,not-valid")


def test_validate_image_value_accepts_valid_image(png_data_uri: str):
    validate_image_value(png_data_uri)


def test_render_value_supports_number_and_boolean():
    number_column = ColumnSettings("amount", "Amount", ColumnType.NUMBER)
    boolean_column = ColumnSettings(
        "completed",
        "Completed",
        ColumnType.BOOLEAN,
        true_text="Done",
        false_text="Pending",
    )

    assert render_value(12.0, number_column) == "12"
    assert render_value(False, boolean_column) == "Pending"


def test_render_template_uses_row_values():
    row = pd.Series({"name": "Alice", "id": "42"})
    assert render_template("Hello {name}", row, {"safe_id": "42-safe"}) == "Hello Alice"


def test_create_single_doc_writes_heading_intro_and_fields(
    tmp_path: Path, png_data_uri: str
):
    output_path = tmp_path / "participant.docx"
    row = pd.Series({"id": "001", "name": "Alice", "sign": png_data_uri})
    columns = [
        ColumnSettings("name", "Name", ColumnType.STRING),
        ColumnSettings("sign", "Signature", ColumnType.IMAGE),
    ]

    create_single_doc(
        row,
        output_path,
        columns,
        DocumentSettings(id_col="id", title="Report", image_width=1, image_height=1),
        intro_text="Welcome Alice",
    )

    doc = Document(output_path)
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert output_path.exists()
    assert texts[0] == "Report"
    assert "Welcome Alice" in texts
    assert any("Name: Alice" in text for text in texts)
    assert len(doc.inline_shapes) == 1


def test_create_doc_with_table_supports_mixed_types(tmp_path: Path, png_data_uri: str):
    output_path = tmp_path / "summary.docx"
    dataframe = pd.DataFrame(
        {
            "id": ["1"],
            "name": ["Alice"],
            "amount": [30],
            "completed": [True],
            "sign": [png_data_uri],
        }
    )
    columns = [
        ColumnSettings("name", "Name", ColumnType.STRING),
        ColumnSettings("amount", "Amount", ColumnType.NUMBER),
        ColumnSettings("completed", "Completed", ColumnType.BOOLEAN),
        ColumnSettings("sign", "Signature", ColumnType.IMAGE),
    ]

    create_doc_with_table(
        dataframe,
        output_path,
        columns,
        DocumentSettings(
            id_col="id", title="Summary", table_image_width=1, table_image_height=1
        ),
    )

    doc = Document(output_path)
    assert output_path.exists()
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text == "Alice"
    assert doc.tables[0].rows[1].cells[1].text == "30"


def test_create_single_doc_falls_back_on_bad_image(tmp_path: Path):
    output_path = tmp_path / "participant.docx"
    row = pd.Series({"id": "001", "sign": "data:image/png;base64,INVALID"})

    create_single_doc(
        row,
        output_path,
        [ColumnSettings("sign", "Signature", ColumnType.IMAGE)],
        DocumentSettings(id_col="id", title="Report"),
    )

    doc = Document(output_path)
    assert any("INVALID" in paragraph.text for paragraph in doc.paragraphs)
