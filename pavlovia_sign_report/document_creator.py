"""Create Word documents from configured report data."""

from __future__ import annotations

import base64
import binascii
import logging
import os
import re
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Inches
from PIL import Image

from pavlovia_sign_report.models import ColumnSettings, ColumnType, DocumentSettings

logger = logging.getLogger(__name__)

_DATA_URI_RE = re.compile(r"^data:image/[^;]+;base64,", re.IGNORECASE)
_BOOL_TRUE_VALUES = {True, "1", "true", "True", "yes", "Yes", "y", "Y"}
_BOOL_FALSE_VALUES = {False, "0", "false", "False", "no", "No", "n", "N"}


class SafeFormatDict(dict[str, str]):
    """Mapping that raises a descriptive error for missing template keys."""

    def __missing__(self, key: str) -> str:
        raise ValueError(f"Unknown placeholder '{key}' in the configured template.")


def _decode_image(data_uri: str) -> BytesIO:
    """Decode a base64 image string or data URI."""

    if not isinstance(data_uri, str):
        raise ValueError("Image values must be strings containing base64 data.")

    raw = _DATA_URI_RE.sub("", data_uri.strip())
    if not raw:
        raise ValueError("Image data is empty.")

    try:
        return BytesIO(base64.b64decode(raw, validate=True))
    except (binascii.Error, ValueError) as exc:
        raise ValueError("Image data is not valid base64.") from exc


def validate_image_value(value: str) -> None:
    """Validate that an image value can be decoded and opened."""

    payload = _decode_image(value)
    try:
        with Image.open(payload) as image:
            image.verify()
    except Exception as exc:  # Pillow raises multiple concrete exceptions.
        raise ValueError(
            "Image data could not be parsed as a supported image file."
        ) from exc


def render_value(value: Any, column: ColumnSettings) -> str:
    """Render a configured cell value as text."""

    if column.column_type == ColumnType.STRING:
        if isinstance(value, (list, dict, set, tuple)):
            raise ValueError(
                f"Column '{column.name}' contains a structured value that "
                "cannot be rendered as text."
            )
        return str(value)

    if column.column_type == ColumnType.NUMBER:
        if isinstance(value, bool):
            raise ValueError(
                f"Column '{column.name}' expects a number, got a boolean value."
            )
        try:
            number = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Column '{column.name}' expects a numeric value."
            ) from exc
        if number.is_integer():
            return str(int(number))
        return str(number)

    if column.column_type == ColumnType.BOOLEAN:
        if value in _BOOL_TRUE_VALUES:
            return column.true_text
        if value in _BOOL_FALSE_VALUES:
            return column.false_text
        raise ValueError(
            f"Column '{column.name}' expects a boolean-like value such as "
            "true/false, yes/no, or 1/0."
        )

    raise ValueError(
        f"Column '{column.name}' uses unsupported text rendering type "
        f"'{column.column_type.value}'."
    )


def render_template(
    template: str, row: pd.Series, extra_context: dict[str, str]
) -> str:
    """Render a user-configurable template using row data."""

    values = SafeFormatDict(extra_context)
    for key, value in row.items():
        values[str(key)] = "" if pd.isna(value) else str(value)
    return template.format_map(values)


def create_single_doc(
    row: pd.Series,
    output_path: str | Path,
    columns: list[ColumnSettings],
    document_settings: DocumentSettings,
    intro_text: str | None = None,
) -> None:
    """Create one participant document."""

    doc = Document()
    doc.add_heading(document_settings.title, 0)
    if intro_text:
        doc.add_paragraph(intro_text)

    for column in columns:
        if column.column_type == ColumnType.IMAGE:
            doc.add_paragraph(f"{column.display_name}:")
            _add_image_to_document(
                doc,
                row[column.name],
                width=document_settings.image_width,
                height=document_settings.image_height,
                column_name=column.name,
                fallback_prefix=f"{column.display_name}: ",
            )
            continue

        value = render_value(row[column.name], column)
        doc.add_paragraph(f"{column.display_name}: {value}")

    doc.save(str(output_path))


def create_doc_with_table(
    df: pd.DataFrame,
    output_path: str | Path,
    columns: list[ColumnSettings],
    document_settings: DocumentSettings,
) -> None:
    """Create a summary document with a table for all rows."""

    doc = Document()
    doc.add_heading(document_settings.title, 0)
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for index, column in enumerate(columns):
        header_cells[index].text = column.display_name or column.name

    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for index, column in enumerate(columns):
            if column.column_type == ColumnType.IMAGE:
                _add_image_to_cell(
                    row_cells[index],
                    row[column.name],
                    width=document_settings.table_image_width,
                    height=document_settings.table_image_height,
                    column_name=column.name,
                )
                continue
            row_cells[index].text = render_value(row[column.name], column)

    doc.save(str(output_path))


def _add_image_to_document(
    doc: Document,
    image_value: str,
    *,
    width: float,
    height: float,
    column_name: str,
    fallback_prefix: str,
) -> None:
    try:
        tmp_path = _write_temp_image(image_value)
        try:
            doc.add_picture(tmp_path, width=Inches(width), height=Inches(height))
        finally:
            _remove_temp_file(tmp_path)
    except Exception:
        logger.error(
            "Failed to embed image for column '%s'", column_name, exc_info=True
        )
        doc.add_paragraph(f"{fallback_prefix}{image_value}")


def _add_image_to_cell(
    cell, image_value: str, *, width: float, height: float, column_name: str
) -> None:
    try:
        tmp_path = _write_temp_image(image_value)
        try:
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run()
            run.add_picture(tmp_path, width=Inches(width), height=Inches(height))
        finally:
            _remove_temp_file(tmp_path)
    except Exception:
        logger.error(
            "Failed to embed image for column '%s'", column_name, exc_info=True
        )
        cell.text = str(image_value)


def _write_temp_image(image_value: str) -> str:
    payload = _decode_image(image_value)
    with Image.open(payload) as image:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp_path = tmp.name
        image.save(tmp_path)
    return tmp_path


def _remove_temp_file(tmp_path: str) -> None:
    try:
        os.remove(tmp_path)
    except OSError:
        logger.warning(
            "Failed to remove temporary image file '%s'", tmp_path, exc_info=True
        )
